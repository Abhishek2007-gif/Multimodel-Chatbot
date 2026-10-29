import io
import os
import difflib
from typing import List, Dict, Any, Tuple, Optional
from app.services.providers import get_llm_provider


class DocumentService:
    @staticmethod
    def extract_text_from_file(file_path: str, mime_type: str, file_name: str) -> str:
        """
        Extract clean text content from PDF, DOCX, TXT, or Markdown files.
        """
        ext = os.path.splitext(file_name)[1].lower()

        if ext == ".pdf" or "pdf" in mime_type:
            try:
                from pypdf import PdfReader
                reader = PdfReader(file_path)
                pages_text = []
                for i, page in enumerate(reader.pages):
                    t = page.extract_text()
                    if t:
                        pages_text.append(f"--- Page {i + 1} ---\n{t}")
                extracted = "\n\n".join(pages_text).strip()
                if not extracted:
                    raise ValueError("No extractable text found in PDF (it may be scanned or image-only).")
                return extracted
            except Exception as e:
                raise ValueError(f"Failed to read PDF document: {str(e)}")

        elif ext in [".docx", ".doc"] or "wordprocessingml" in mime_type:
            try:
                from docx import Document
                doc = Document(file_path)
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                # Also include table text
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if row_text:
                            paragraphs.append(row_text)
                extracted = "\n\n".join(paragraphs).strip()
                if not extracted:
                    raise ValueError("No text found in DOCX document.")
                return extracted
            except Exception as e:
                raise ValueError(f"Failed to read DOCX document: {str(e)}")

        elif ext in [".txt", ".md", ".markdown", ".csv", ".json", ".py", ".js", ".html"]:
            try:
                with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                    return f.read()
            except Exception as e:
                raise ValueError(f"Failed to read text document: {str(e)}")

        else:
            raise ValueError(f"Unsupported document format '{ext}'. Supported formats: PDF, DOCX, TXT, Markdown.")

    @classmethod
    async def create_document(cls, prompt: str, context: Optional[str] = None) -> str:
        """Generate a structured document draft from a user prompt."""
        provider = get_llm_provider()
        system_prompt = (
            "You are an expert technical and professional document writer. "
            "Draft a comprehensive, well-structured document based on the user's prompt. "
            "Use clear Markdown formatting with titles, headings (#, ##), bullet points, and paragraphs."
        )
        user_content = prompt
        if context:
            user_content += f"\n\nReference context:\n{context}"

        messages = [{"role": "user", "content": user_content}]
        return await provider.generate(messages, system_prompt=system_prompt, temperature=0.7, max_tokens=2500)

    @classmethod
    async def edit_document(
        cls,
        original_content: str,
        instruction: str,
        action_type: str = "custom",
        target_language: str = "English"
    ) -> str:
        """
        Apply requested edits to the document, strictly preserving unedited sections.
        """
        provider = get_llm_provider()

        if action_type == "summarize":
            task_desc = "Provide an executive summary of the document while retaining key metrics, decisions, and takeaways."
        elif action_type == "grammar":
            task_desc = "Fix spelling, grammar, punctuation, and improve phrasing while preserving the original meaning and structure."
        elif action_type == "rewrite":
            task_desc = "Rewrite and polish the document for enhanced clarity, professional tone, and flow."
        elif action_type == "translate":
            task_desc = f"Translate the complete document accurately into {target_language}, preserving formatting."
        else:
            task_desc = f"Apply this specific revision: {instruction}. Preserve existing sections where no change was requested."

        system_prompt = f"""You are an expert document editor.
Task: {task_desc}

Rules:
1. Preserve unchanged paragraphs and document structure wherever possible.
2. Return ONLY the revised document text in Markdown format.
3. Do not include conversational remarks, meta-commentary, or preamble like 'Here is your revised document'.
"""
        messages = [
            {"role": "user", "content": f"Original Document:\n\n{original_content}\n\nRevision Instruction:\n{instruction}"}
        ]
        revised = await provider.generate(messages, system_prompt=system_prompt, temperature=0.3, max_tokens=3000)
        return revised.strip()

    @staticmethod
    def compute_diff(original: str, revised: str) -> List[Dict[str, str]]:
        """
        Compute line-by-line and chunk-level diff between original and revised text.
        Returns a list of diff chunks with types: 'equal', 'insert', 'delete'.
        """
        matcher = difflib.SequenceMatcher(None, original.splitlines(keepends=True), revised.splitlines(keepends=True))
        diff_chunks = []

        for tag, i1, i2, j1, j2 in matcher.get_opcodes():
            orig_slice = "".join(original.splitlines(keepends=True)[i1:i2])
            rev_slice = "".join(revised.splitlines(keepends=True)[j1:j2])

            if tag == "equal":
                diff_chunks.append({"type": "equal", "text": orig_slice})
            elif tag == "delete":
                diff_chunks.append({"type": "delete", "text": orig_slice})
            elif tag == "insert":
                diff_chunks.append({"type": "insert", "text": rev_slice})
            elif tag == "replace":
                diff_chunks.append({"type": "delete", "text": orig_slice})
                diff_chunks.append({"type": "insert", "text": rev_slice})

        return diff_chunks

    @staticmethod
    def export_document(content: str, title: str, format_type: str) -> Tuple[bytes, str, str]:
        """
        Export content as downloadable bytes for DOCX, PDF, MD, or TXT.
        Returns (file_bytes, mime_type, file_extension).
        """
        format_type = format_type.lower().strip()

        if format_type in ["docx", "doc"]:
            from docx import Document
            doc = Document()
            doc.add_heading(title or "Document", level=1)
            for paragraph in content.split("\n\n"):
                p_text = paragraph.strip()
                if p_text.startswith("# "):
                    doc.add_heading(p_text[2:], level=1)
                elif p_text.startswith("## "):
                    doc.add_heading(p_text[3:], level=2)
                elif p_text.startswith("### "):
                    doc.add_heading(p_text[4:], level=3)
                elif p_text.startswith("- ") or p_text.startswith("* "):
                    for line in p_text.split("\n"):
                        if line.startswith("- ") or line.startswith("* "):
                            doc.add_paragraph(line[2:], style='List Bullet')
                        else:
                            doc.add_paragraph(line)
                elif p_text:
                    doc.add_paragraph(p_text)

            bio = io.BytesIO()
            doc.save(bio)
            return bio.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"

        elif format_type == "pdf":
            from reportlab.lib.pagesizes import letter
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            import html

            bio = io.BytesIO()
            doc = SimpleDocTemplate(bio, pagesize=letter, leftMargin=54, rightMargin=54, topMargin=54, bottomMargin=54)
            styles = getSampleStyleSheet()
            flowables = []

            # Title
            title_style = ParagraphStyle(
                'DocTitle',
                parent=styles['Heading1'],
                fontSize=20,
                leading=24,
                spaceAfter=14
            )
            flowables.append(Paragraph(html.escape(title or "Document"), title_style))
            flowables.append(Spacer(1, 10))

            body_style = styles['Normal']
            body_style.fontSize = 11
            body_style.leading = 15

            for block in content.split("\n\n"):
                block = block.strip()
                if not block:
                    continue
                safe_text = html.escape(block).replace("\n", "<br/>")
                flowables.append(Paragraph(safe_text, body_style))
                flowables.append(Spacer(1, 10))

            doc.build(flowables)
            return bio.getvalue(), "application/pdf", "pdf"

        elif format_type in ["md", "markdown"]:
            return content.encode("utf-8"), "text/markdown; charset=utf-8", "md"

        else:  # plain text default
            return content.encode("utf-8"), "text/plain; charset=utf-8", "txt"
